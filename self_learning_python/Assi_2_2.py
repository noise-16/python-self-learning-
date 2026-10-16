from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

OUTPUT_FILE = "AIML_Batch_Grades.xlsx"
PASS_MARK = 40

# 1. Create workbook and name the active sheet
wb = Workbook()
ws = wb.active
ws.title = "Semester_Results"

# 2. Header row
headers = [
    "Roll No",
    "Student Name",
    "Lab Marks (out of 50)",
    "Theory Marks (out of 50)",
    "Total Marks (out of 100)",
    "Result",
]
ws.append(headers)

# 3. Sample student data: (Roll No, Name, Lab, Theory)
students = [
    ("AIML001", "Aarav Sharma", 44, 41),
    ("AIML002", "Diya Mehta", 38, 35),
    ("AIML003", "Rohan Iyer", 25, 12),   # expected Fail
    ("AIML004", "Sneha Kulkarni", 47, 45),
    ("AIML005", "Kabir Singh", 30, 28),
    ("AIML006", "Ananya Desai", 18, 15),  # expected Fail (33 total)
]
for roll, name, lab, theory in students:
    ws.append([roll, name, lab, theory, None, None])

first_data_row = 2
last_data_row = ws.max_row

# 4. Iterate over data rows: compute Total and Result.
#    Written as live formulas so the sheet recalculates if marks are edited.
for row in range(first_data_row, last_data_row + 1):
    ws[f"E{row}"] = f"=C{row}+D{row}"
    ws[f"F{row}"] = f'=IF(E{row}>={PASS_MARK},"Pass","Fail")'

# 5. Summary row: class averages
summary_row = last_data_row + 1
ws[f"A{summary_row}"] = "Class Average"
for col in ("C", "D", "E"):
    ws[f"{col}{summary_row}"] = (
        f"=ROUND(AVERAGE({col}{first_data_row}:{col}{last_data_row}),2)"
    )

# Styling
base_font = Font(name="Arial", size=11)
bold_font = Font(name="Arial", size=11, bold=True)
header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")

for row in ws.iter_rows(min_row=1, max_row=summary_row, max_col=len(headers)):
    for cell in row:
        cell.font = base_font
        if cell.column >= 3:
            cell.alignment = Alignment(horizontal="center")

for cell in ws[1]:                       # bold + light-gray header
    cell.font = bold_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

for cell in ws[summary_row]:             # bold summary row
    cell.font = bold_font
    cell.fill = header_fill

# Column widths
for idx, width in enumerate([12, 22, 22, 24, 24, 10], start=1):
    ws.column_dimensions[get_column_letter(idx)].width = width
ws.freeze_panes = "A2"

wb.save(OUTPUT_FILE)
print(f"Saved {OUTPUT_FILE}")